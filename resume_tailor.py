# mypy: ignore-errors
import json
import re
import tempfile
import time
from pathlib import Path

import requests
from docx import Document
from docx2pdf import convert

from config import GROQ_API_KEY, GROQ_MODEL
from logging_config import logger

BASE_RESUME = Path(__file__).with_name("resume_template.docx")


def _read_base_resume():
    doc = Document(str(BASE_RESUME))
    return "\n".join(p.text for p in doc.paragraphs if p.text.strip())


def _extract_json(text):
    text = text.strip()
    if text.startswith("```"):
        text = re.sub(r"^```(?:json)?\s*", "", text)
        text = re.sub(r"\s*```$", "", text)
    return json.loads(text)


MAX_JD_CHARS = 8000
MAX_BASE_RESUME_CHARS = 12000
MAX_TOTAL_PROMPT_CHARS = 22000
GROQ_MAX_RETRIES = 3


def _clean_job_description(text):
    text = text or ""
    # Remove common recommendation/footer noise from job-alert emails.
    noise_patterns = [
        r"\b\d+\s+more\s+[^\n]{0,180}jobs?[^\n]*",
        r"(?im)^\s*(easily apply|actively hiring|apply now|view job|see all jobs)\s*$",
        r"(?im)^\s*(unsubscribe|manage preferences|privacy policy|terms of use)\s*$",
    ]
    for pattern in noise_patterns:
        text = re.sub(pattern, " ", text)
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()[:MAX_JD_CHARS]


def _post_groq(payload):
    last_error = None
    for attempt in range(GROQ_MAX_RETRIES):
        try:
            response = requests.post(
                "https://api.groq.com/openai/v1/chat/completions",
                headers={
                    "Authorization": f"Bearer {GROQ_API_KEY}",
                    "Content-Type": "application/json",
                },
                json=payload,
                timeout=90,
            )
            if response.status_code == 429:
                wait = 3 * (2**attempt)
                logger.warning(
                    "Groq rate limit (429). Retrying in %ds (%d/%d)...",
                    wait,
                    attempt + 1,
                    GROQ_MAX_RETRIES,
                )
                time.sleep(wait)
                last_error = requests.HTTPError("Groq rate limit", response=response)
                continue
            if response.status_code == 413:
                raise requests.HTTPError(
                    "Groq payload too large after prompt limits.", response=response
                )
            response.raise_for_status()
            return response
        except requests.RequestException as exc:
            last_error = exc
            if (
                attempt < GROQ_MAX_RETRIES - 1
                and getattr(exc, "response", None) is not None
                and exc.response.status_code >= 500
            ):
                time.sleep(2 * (attempt + 1))
                continue
            raise
    raise last_error


def tailor_resume(job):
    """Create a JD-specific ATS resume without inventing experience."""
    if not GROQ_API_KEY:
        logger.warning("GROQ_API_KEY missing; cannot generate tailored resume.")
        return None

    base = _read_base_resume()
    jd = _clean_job_description(job.get("description", ""))
    title = job.get("title", "AI/ML Engineer")
    company = job.get("company", "Unknown")

    base = base[:MAX_BASE_RESUME_CHARS]

    prompt_template = """
You are an ATS resume editor. Tailor the candidate resume below to the exact job description.

STRICT RULES:
1. NEVER invent employers, dates, degrees, certifications, tools, projects, metrics, or responsibilities.
2. Only use facts explicitly present in the base resume.
3. You may reorder existing skills and rewrite bullets to emphasize facts relevant to the JD.
4. Do not claim a tool/skill merely because it is similar to another tool.
5. Keep the candidate's name/contact information unchanged.
6. Keep all employment dates and employer names unchanged.
7. Make the resume ATS-friendly and concise, ideally 1-2 pages.
8. Use exact JD terminology only when it accurately describes an existing skill.
9. If a JD requirement is missing from the base resume, do not add it.
10. Return JSON only.

JOB TITLE: {title}
COMPANY: {company}

JOB DESCRIPTION:
{jd}

BASE RESUME:
{base}

Return exactly this JSON structure:
{{
  "headline": "...",
  "summary": "...",
  "skills": ["category: skill1, skill2"],
  "experience": [
    {{"role": "...", "company": "...", "dates": "...", "bullets": ["...", "..."]}}
  ],
  "projects": [
    {{"name": "...", "stack": "...", "bullets": ["..."]}}
  ],
  "education": ["..."]
}}
"""

    # Keep the request comfortably below Groq payload/context limits.
    fixed = prompt_template.format(title=title, company=company, jd="", base="")
    available = max(4000, MAX_TOTAL_PROMPT_CHARS - len(fixed))
    jd_budget = min(MAX_JD_CHARS, int(available * 0.45))
    base_budget = min(MAX_BASE_RESUME_CHARS, available - jd_budget)
    jd = jd[:jd_budget]
    base = base[:base_budget]
    prompt = prompt_template.format(title=title, company=company, jd=jd, base=base)
    if len(prompt) > MAX_TOTAL_PROMPT_CHARS:
        # Last-resort deterministic trim of the JD only; preserve the JSON instructions.
        overflow = len(prompt) - MAX_TOTAL_PROMPT_CHARS
        jd = jd[: max(3000, len(jd) - overflow - 100)]
        prompt = prompt_template.format(title=title, company=company, jd=jd, base=base)

    payload = {
        "model": GROQ_MODEL,
        "temperature": 0.1,
        "response_format": {"type": "json_object"},
        "messages": [
            {"role": "system", "content": "Return valid JSON only."},
            {"role": "user", "content": prompt},
        ],
    }
    response = _post_groq(payload)
    content = response.json()["choices"][0]["message"]["content"]
    return _extract_json(content)


def build_resume_docx(job, tailored):
    safe_company = re.sub(r"[^A-Za-z0-9]+", "_", job.get("company", "Company")).strip("_")
    safe_title = re.sub(r"[^A-Za-z0-9]+", "_", job.get("title", "Role")).strip("_")
    filename = f"Viral_Sherathiya_{safe_company}_{safe_title}_Tailored_Resume.docx"[:180]
    path = Path(tempfile.gettempdir()) / filename

    doc = Document()
    styles = doc.styles
    styles["Normal"].font.name = "Arial"
    styles["Normal"].font.size = __import__("docx").shared.Pt(9)

    p = doc.add_paragraph()
    p.alignment = 1
    r = p.add_run("VIRAL SHERATHIYA")
    r.bold = True
    r.font.size = __import__("docx").shared.Pt(18)

    p = doc.add_paragraph()
    p.alignment = 1
    p.add_run(tailored.get("headline", "AI/ML ENGINEER")).bold = True

    p = doc.add_paragraph()
    p.alignment = 1
    p.add_run(
        "Ahmedabad, Gujarat, India | +91-8140408415 | viralsherathiya1008@gmail.com | linkedin.com/in/viralsherathiya | github.com/Viral3899"
    )

    def heading(text):
        p = doc.add_paragraph()
        r = p.add_run(text.upper())
        r.bold = True
        r.font.size = __import__("docx").shared.Pt(11)

    heading("Professional Summary")
    doc.add_paragraph(tailored.get("summary", ""))

    heading("Core Skills")
    for item in tailored.get("skills", []):
        doc.add_paragraph(item, style="List Bullet")

    heading("Professional Experience")
    for exp in tailored.get("experience", []):
        p = doc.add_paragraph()
        r = p.add_run(f"{exp.get('role','')} — {exp.get('company','')}")
        r.bold = True
        p.add_run(f" | {exp.get('dates','')}")
        for bullet in exp.get("bullets", []):
            doc.add_paragraph(bullet, style="List Bullet")

    heading("Selected Projects")
    for project in tailored.get("projects", []):
        p = doc.add_paragraph()
        r = p.add_run(project.get("name", ""))
        r.bold = True
        if project.get("stack"):
            p.add_run(f" — {project['stack']}")
        for bullet in project.get("bullets", []):
            doc.add_paragraph(bullet, style="List Bullet")

    heading("Education")
    for item in tailored.get("education", []):
        doc.add_paragraph(item, style="List Bullet")

    doc.save(str(path))
    return str(path)


def build_resume_pdf(job, tailored):
    """Generate PDF resume from DOCX."""
    docx_path = build_resume_docx(job, tailored)
    pdf_path = str(Path(docx_path).with_suffix(".pdf"))
    try:
        convert(docx_path, pdf_path)
        logger.info("PDF resume generated: %s", pdf_path)
        return pdf_path
    except Exception as exc:
        logger.error("Failed to convert DOCX to PDF: %s", exc)
        # Fallback to docx if PDF conversion fails
        return docx_path
